"""Юнит-тесты отчётов §9: Excel, JSON, HTML.

Ловушки, закреплённые здесь (§11.1):

* текст с ведущим ``=`` в Excel экранируется, иначе ячейка становится
  формулой и даёт ``#VALUE!`` (§9.1);
* ``Decimal`` в JSON уходит **строкой**: ``float`` теряет копейки и ломает
  побитовое сравнение эталона §11.3;
* повторный дамп одного результата даёт **побитово идентичный** JSON (§13.5);
* ни один отчёт не квалифицирует находку как хищение (§1.2, §13.10).
"""

from __future__ import annotations

import json
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from cashforensics.models import (
    AnalysisResult,
    BalanceTrace,
    CausalLine,
    FileMeta,
    Finding,
    FindingCode,
    Materiality,
    NonNegativityReport,
    Severity,
    ValidationReport,
    Waterfall,
)
from cashforensics.report.excel import SHEETS, escape_leading_equals, render_excel, write_row
from cashforensics.report.html import render_html, render_section_limitations
from cashforensics.report.json_out import default_encoder, render_json

pytestmark = pytest.mark.unit

ZERO = Decimal("0.00")
DAY = date(2026, 5, 31)

FORBIDDEN_WORDS = ("фрод", "кража", "виновен", "хищение", "мошенничеств")
"""§1.2, §13.10: инструмент не квалифицирует находки. Умысел из данных не выводим."""


@pytest.fixture
def result() -> AnalysisResult:
    """Минимальный результат анализа для проверки отчётов."""
    finding = Finding(
        code=FindingCode.RKO_OVERSTATED,
        severity=Severity.ERROR,
        date=DAY,
        amount=Decimal("160.89"),
        ledger_rows=[987],
        doc_numbers=["00000257506"],
        title="Проводка возврата больше фактической выдачи",
        explanation="В 1С 289.31, в опер-логе 128.42. Требует проверки.",
        confidence=1.0,
        materiality=Materiality.IMMATERIAL,
        balance_impact=Decimal("160.89"),
    )
    waterfall = Waterfall(
        closing=Decimal("-160.89"),
        target=ZERO,
        income_delta=ZERO,
        collection_delta=ZERO,
        refund_delta=Decimal("-160.89"),
        other_delta=ZERO,
        log_imbalance=ZERO,
        unresolved=ZERO,
        causal_lines=(
            CausalLine(
                amount=Decimal("160.89"),
                title="=РКО 00000257506 больше выдачи",
                finding_index=0,
                doc_numbers=("00000257506",),
            ),
        ),
    )
    return AnalysisResult(
        meta=FileMeta(
            path=Path("PAX_119049688.xlsx"),
            sha256="0" * 64,
            pvz_id="PAX 119049688",
            stated_period=None,
            actual_period=(date(2026, 1, 1), DAY),
            rules_version="1.0",
            config_sha256="1" * 64,
            run_timestamp=datetime(2026, 8, 7, 12, 0, 0),
        ),
        validation=ValidationReport(
            checks=(),
            hard_failed=False,
            daily_reconciliation_enabled=True,
            unknown_accounts=(),
            cutoff_suspected=False,
        ),
        reconciliation={},
        timing={},
        balance_trace=BalanceTrace(
            opening=ZERO,
            closing_computed=Decimal("-160.89"),
            target=ZERO,
            nonnegativity=NonNegativityReport(
                first_below_target=DAY,
                never_recovered_after=DAY,
                minimum=Decimal("-160.89"),
                minimum_date=DAY,
                share_of_days_below=0.5,
            ),
            sweep_points=(),
            change_points=(),
            shift_windows=(),
            quarterly_balances=(),
        ),
        waterfall=waterfall,
        findings=[finding],
        signatures=[],
        localizations=[],
        unresolved=ZERO,
    )


class TestEscapeLeadingEquals:
    """§9.1 — ловушка «текст стал формулой»."""

    def test_leading_equals_is_escaped(self) -> None:
        assert escape_leading_equals("=РКО 001") == "'=РКО 001"

    def test_ordinary_text_is_untouched(self) -> None:
        assert escape_leading_equals("РКО 001") == "РКО 001"

    def test_equals_inside_is_untouched(self) -> None:
        assert escape_leading_equals("сумма = 100") == "сумма = 100"

    def test_written_cell_is_text_not_formula(self, tmp_path: Path) -> None:
        """Проверка на настоящем файле: ячейка не должна быть формулой."""
        workbook = Workbook()
        sheet = workbook.worksheets[0]
        write_row(sheet, ["=РКО 00000257506", Decimal("160.89")])
        path = tmp_path / "проверка.xlsx"
        workbook.save(path)
        workbook.close()

        reopened = load_workbook(path)
        cell = reopened.worksheets[0].cell(row=1, column=1)
        assert cell.data_type == "s"
        assert cell.value == "'=РКО 00000257506"
        reopened.close()

    def test_decimal_stays_a_number(self, tmp_path: Path) -> None:
        """§9.1: сумма пишется числом, а не строкой — иначе формат бесполезен.

        Читается она обратно уже как ``float``: числовая ячейка `.xlsx` — это
        IEEE-754, и граница проходит по чтению, а не по записи. Поэтому
        проверяются оба конца: тип ячейки числовой, а в самом файле лежит
        точный десятичный текст без расширения до двоичной дроби.
        """
        workbook = Workbook()
        write_row(workbook.worksheets[0], [Decimal("160.89")])
        path = tmp_path / "сумма.xlsx"
        workbook.save(path)
        workbook.close()

        reopened = load_workbook(path)
        cell = reopened.worksheets[0].cell(row=1, column=1)
        assert cell.data_type == "n"
        assert cell.value == pytest.approx(160.89)
        reopened.close()

        with zipfile.ZipFile(path) as archive:
            sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert "<v>160.89</v>" in sheet_xml


class TestExcelReport:
    """§9.1 — состав и формулировки листов."""

    def test_all_required_sheets_are_present(
        self,
        result: AnalysisResult,
        tmp_path: Path,
    ) -> None:
        path = render_excel(result, tmp_path / "отчёт.xlsx")
        workbook = load_workbook(path)
        assert workbook.sheetnames == list(SHEETS)
        workbook.close()

    def test_audit_sheet_only_on_demand(self, result: AnalysisResult, tmp_path: Path) -> None:
        """§12: журнал автогашения — по флагу ``--audit-log``."""
        path = render_excel(result, tmp_path / "аудит.xlsx", audit_log=True)
        workbook = load_workbook(path)
        assert len(workbook.sheetnames) == len(SHEETS) + 1
        workbook.close()

    def test_causal_line_with_leading_equals_survives(
        self,
        result: AnalysisResult,
        tmp_path: Path,
    ) -> None:
        """Ловушка §9.1 на настоящем отчёте, а не на изолированной функции."""
        path = render_excel(result, tmp_path / "причина.xlsx")
        workbook = load_workbook(path)
        values = [
            cell.value
            for row in workbook[SHEETS[0]].iter_rows()
            for cell in row
            if isinstance(cell.value, str)
        ]
        assert any(value.startswith("'=") for value in values)
        assert not any(value.startswith("=") for value in values)
        workbook.close()

    def test_wording_does_not_qualify_the_finding(
        self,
        result: AnalysisResult,
        tmp_path: Path,
    ) -> None:
        """§1.2, §13.10: «расхождение», но не «фрод» и не «кража»."""
        path = render_excel(result, tmp_path / "формулировки.xlsx")
        workbook = load_workbook(path)
        text = " ".join(
            str(cell.value).lower()
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        for word in FORBIDDEN_WORDS:
            assert word not in text
        workbook.close()


class TestJsonReport:
    """§9.2 — сериализация и §13.5 побитовая воспроизводимость."""

    def test_decimal_is_a_string(self) -> None:
        assert default_encoder(Decimal("160.89")) == "160.89"

    def test_date_is_iso(self) -> None:
        assert default_encoder(DAY) == "2026-05-31"

    def test_unknown_type_raises_instead_of_str(self) -> None:
        """Молчаливая подмена на ``str`` скрыла бы дефект модели."""
        with pytest.raises(TypeError, match="не сериализуем"):
            default_encoder(object())

    def test_amounts_keep_their_kopecks(self, result: AnalysisResult, tmp_path: Path) -> None:
        payload = json.loads(render_json(result, tmp_path / "о.json").read_text(encoding="utf-8"))
        assert payload["findings"][0]["amount"] == "160.89"
        assert isinstance(payload["findings"][0]["amount"], str)

    def test_repeat_dump_is_byte_identical(
        self,
        result: AnalysisResult,
        tmp_path: Path,
    ) -> None:
        """§13.5: один вход при одной версии правил — побитово один выход."""
        first = render_json(result, tmp_path / "первый.json").read_bytes()
        second = render_json(result, tmp_path / "второй.json").read_bytes()
        assert first == second

    def test_keys_are_sorted(self, result: AnalysisResult, tmp_path: Path) -> None:
        """Без ``sort_keys`` порядок зависел бы от объявления полей."""
        text = render_json(result, tmp_path / "ключи.json").read_text(encoding="utf-8")
        payload = json.loads(text)
        assert list(payload) == sorted(payload)


class TestHtmlReport:
    """§9.3 — опциональный одностраничный отчёт."""

    def test_limitations_section_is_always_present(self, result: AnalysisResult) -> None:
        """§16: раздел ограничений обязателен в каждом отчёте."""
        section = render_section_limitations(result)
        assert "какая из двух систем права" in section
        assert "Дисбаланс самих логов" in section

    def test_finding_text_is_escaped(self, result: AnalysisResult, tmp_path: Path) -> None:
        payload = "<script>alert(1)</script>"
        finding = result.findings[0].model_copy(update={"explanation": payload})
        injected = result.model_copy(update={"findings": [finding]})
        document = render_html(injected, tmp_path / "о.html").read_text(encoding="utf-8")
        assert "<script>alert(1)</script>" not in document
        assert "&lt;script&gt;" in document

    def test_wording_does_not_qualify_the_finding(
        self,
        result: AnalysisResult,
        tmp_path: Path,
    ) -> None:
        document = render_html(result, tmp_path / "формулировки.html").read_text(encoding="utf-8")
        for word in FORBIDDEN_WORDS:
            assert word not in document.lower()
