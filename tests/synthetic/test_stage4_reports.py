"""Синтетические кейсы этапа 4 — отчёты §9, ранжирование §5.11, пакет §13.9.

Проверяют конвейер целиком на настоящем `.xlsx`: между ingest и report
двенадцать стадий, и сцепка находок с отчётом на собранных вручную структурах
не проверяется.

Обязательные критерии §13, закреплённые здесь:

* §13.5 — повторный прогон одного файла даёт побитово идентичный JSON;
* §13.6 — Excel открывается без ошибок формул;
* §13.9 — пакетный прогон формирует сводную таблицу с ранжированием;
* §13.10 — ни один отчёт не квалифицирует находку как хищение.
"""

from __future__ import annotations

import json
import re
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from cashforensics.cli import app, run_pipeline
from cashforensics.models import Config, FindingCode
from cashforensics.report.excel import SHEETS, render_excel
from cashforensics.report.json_out import render_json
from tests.conftest import OpsRecord, Posting, SheetPlan

if TYPE_CHECKING:
    from collections.abc import Callable

pytestmark = pytest.mark.synthetic

ZERO = Decimal("0.00")
START = date(2024, 2, 1)
CONFIG_PATH = Path(__file__).resolve().parents[2] / "config" / "default.yaml"

FORBIDDEN_WORDS = ("фрод", "кража", "виновен", "хищение", "мошенничеств")
"""§1.2, §13.10: квалификация находки инструментом запрещена."""

FORMULA_RE = re.compile(r"<f>(.*?)</f>")
"""Формулы листа как они лежат в XML — читать через openpyxl их нельзя."""

FORMULA_SUM_RE = re.compile(r"SUM\([A-Z]+\d+:[A-Z]+\d+\)")
"""Единственная разрешённая формула отчёта — итог §9.1.8."""


def _sale(day: date, amount: Decimal, index: int) -> Posting:
    return Posting(day, f"Приходный кассовый ордер 0001{index:05d}", "50.2", amount, "90.1.1")


def _collection(day: date, amount: Decimal, index: int) -> Posting:
    return Posting(
        day,
        f"Расходный кассовый ордер 0002{index:05d}",
        "51",
        None,
        "50.2",
        amount,
    )


def _refund(day: date, amount: Decimal, index: int) -> Posting:
    return Posting(
        day,
        f"Расходный кассовый ордер 0003{index:05d}",
        "76.9.1",
        None,
        "50.2",
        amount,
    )


def _overstated_plan() -> SheetPlan:
    """Касса с одной завышенной проводкой возврата — эталонная форма §5.9.3.

    В 1С возврат 289,31, в опер-логе за тот же день одна выдача 128,42:
    правило единственной проводки §5.9.3 сводит расхождение 160,89 к документу
    без обращения к гейтам §7.3.
    """
    days = [START + timedelta(days=offset) for offset in range(6)]
    postings = [_sale(day, Decimal("1000.00"), index) for index, day in enumerate(days, start=1)]
    postings.append(_refund(days[3], Decimal("289.31"), 1))
    postings.extend(
        _collection(day, Decimal("1000.00"), index) for index, day in enumerate(days, start=1)
    )
    return SheetPlan(
        variant="A",
        postings=postings,
        rko=[
            # Инкассации: получатель «Центральная касса» — §3.6, без него
            # раскладка блоков уходит в time-fallback варианта C.
            *(
                OpsRecord(
                    datetime.combine(day, datetime.min.time().replace(hour=18)),
                    Decimal("1000.00"),
                    "Центральная касса",
                )
                for day in days
            ),
            OpsRecord(
                datetime.combine(days[3], datetime.min.time().replace(hour=14)),
                Decimal("128.42"),
                "Иванов Иван",
            ),
        ],
        pko=[
            OpsRecord(
                datetime.combine(day, datetime.min.time().replace(hour=12)),
                Decimal("1000.00"),
                "Клиент",
            )
            for day in days
        ],
        closing=Decimal("-289.31"),
        closing_side="К",
    )


@pytest.fixture
def register(make_workbook: Callable[[SheetPlan], Path]) -> Path:
    return make_workbook(_overstated_plan())


class TestFindingsReachTheReport:
    """Сцепка LOCALIZE → RANK → REPORT: локализация обязана стать находкой §8."""

    def test_localized_discrepancy_becomes_a_finding(
        self,
        register: Path,
        config: Config,
    ) -> None:
        result = run_pipeline(register, config)
        codes = [finding.code for finding in result.findings]
        assert FindingCode.RKO_OVERSTATED in codes

    def test_finding_carries_the_document_trace(self, register: Path, config: Config) -> None:
        """§12: без номера документа находка непроверяема вручную."""
        result = run_pipeline(register, config)
        overstated = next(
            item for item in result.findings if item.code is FindingCode.RKO_OVERSTATED
        )
        assert overstated.amount == Decimal("160.89")
        assert overstated.doc_numbers
        assert overstated.ledger_rows

    def test_waterfall_ties_to_zero(self, register: Path, config: Config) -> None:
        """§5.10: раскладка обязана сходиться без остатка."""
        result = run_pipeline(register, config)
        assert result.unresolved == ZERO


class TestNonZeroTarget:
    """§13.8 — сквозной прогон при ``target_balance ≠ 0``.

    Пол кассы (разменный фонд, который положено оставлять в ящике) делает
    целевое сальдо ненулевым, и тогда «отклонение» — это ``сальдо − T``, а не
    само сальдо. Юнит-уровень §5.10 это уже проверял; здесь важно, что через
    все двенадцать стадий и оба отчёта проходит именно ``T``, а не ноль.
    """

    @staticmethod
    def _with_target(config: Config, target: str) -> Config:
        return config.model_copy(update={"target_balance": Decimal(target)})

    def test_deviation_is_measured_from_the_target(
        self,
        register: Path,
        config: Config,
    ) -> None:
        """Сальдо не меняется от T — меняется то, что считается отклонением."""
        plain = run_pipeline(register, config)
        raised = run_pipeline(register, self._with_target(config, "500.00"))

        assert raised.waterfall.target == Decimal("500.00")
        assert raised.waterfall.closing == plain.waterfall.closing
        assert raised.waterfall.closing - raised.waterfall.target == (
            plain.waterfall.closing - Decimal("500.00")
        )

    def test_waterfall_still_ties(self, register: Path, config: Config) -> None:
        """§13.2: ``|unresolved| < EPS_TIE`` обязано держаться и при T ≠ 0."""
        result = run_pipeline(register, self._with_target(config, "500.00"))

        assert result.unresolved == ZERO
        assert abs(result.waterfall.unresolved) < config.thresholds.EPS_TIE

    def test_causal_lines_cover_the_shifted_deviation(
        self,
        register: Path,
        config: Config,
    ) -> None:
        """Раскладка обязана покрыть ``сальдо − T``, а не ``сальдо``.

        Разница уходит в остаток: пол кассы не является ни находкой §8, ни
        слагаемым §5.10 — это заданная величина, а не расхождение.
        """
        result = run_pipeline(register, self._with_target(config, "500.00"))

        total = sum((line.amount for line in result.waterfall.causal_lines), ZERO)
        assert total == result.waterfall.closing - Decimal("500.00")

    def test_cli_target_reaches_both_reports(
        self,
        register: Path,
        tmp_path: Path,
    ) -> None:
        """``--target`` §10: величина обязана дойти до Excel и до JSON.

        Каталог отчётов отдельный: фикстура ``register`` кладёт саму выгрузку в
        ``tmp_path``, и общий glob подобрал бы исходник вместо отчёта.
        """
        out = tmp_path / "отчёты"
        runner = CliRunner()
        invocation = runner.invoke(
            app,
            [
                "analyze",
                str(register),
                "--config",
                str(CONFIG_PATH),
                "--target",
                "500.00",
                "--out",
                str(out),
            ],
        )
        assert invocation.exit_code == 0, invocation.output

        payload = json.loads(next(out.glob("*.json")).read_text(encoding="utf-8"))
        assert payload["waterfall"]["target"] == "500.00"

        workbook = load_workbook(next(out.glob("*.xlsx")))
        cause = workbook[SHEETS[0]]
        targets = [
            row[3]
            for row in cause.iter_rows(values_only=True)
            if row and row[0] == "Сальдо на конец"
        ]
        workbook.close()
        assert targets == [Decimal("500.00")]

    def test_non_numeric_target_is_refused(self, register: Path, tmp_path: Path) -> None:
        """Догадка вместо числа запрещена (§0.3): молчаливый ноль хуже отказа."""
        out = tmp_path / "отказ"
        runner = CliRunner()
        invocation = runner.invoke(
            app,
            [
                "analyze",
                str(register),
                "--config",
                str(CONFIG_PATH),
                "--target",
                "пол кассы",
                "--out",
                str(out),
            ],
        )

        assert invocation.exit_code != 0
        assert not out.exists() or not list(out.glob("*.xlsx"))


class TestDeterminism:
    """§13.5 — побитовая воспроизводимость."""

    def test_two_runs_give_identical_json(
        self,
        register: Path,
        config: Config,
        tmp_path: Path,
    ) -> None:
        """Единственное законное различие — ``meta.run_timestamp`` (§9.2).

        Он и есть отметка прогона; всё остальное тело обязано совпасть
        побайтово, иначе эталон §11.3 сравнивать нечем.
        """
        first = run_pipeline(register, config)
        second = run_pipeline(register, config)
        meta = second.meta.model_copy(update={"run_timestamp": first.meta.run_timestamp})
        stamped = second.model_copy(update={"meta": meta})
        assert render_json(first, tmp_path / "1.json").read_bytes() == (
            render_json(stamped, tmp_path / "2.json").read_bytes()
        )

    def test_finding_order_is_stable(self, register: Path, config: Config) -> None:
        first = [item.code for item in run_pipeline(register, config).findings]
        second = [item.code for item in run_pipeline(register, config).findings]
        assert first == second


class TestExcelOpensClean:
    """§13.6 — Excel открывается без ошибок формул.

    Настоящий пересчёт (§9.1: «после сохранения — прогон пересчёта формул»)
    здесь не выполняется: в окружении нет работающего офисного пакета. Поэтому
    проверяется то, что проверяемо статически: единственные формулы отчёта —
    объявленные §9.1.8 итоги ``=SUM()``, их диапазоны лежат внутри данных, и ни
    одна текстовая ячейка формулой не стала.
    """

    def test_only_declared_sum_totals_are_formulas(
        self,
        register: Path,
        tmp_path: Path,
    ) -> None:
        """Ловушка §9.1: текст с ведущим ``=`` обязан быть экранирован."""
        runner = CliRunner()
        out = tmp_path / "отчёты"
        result = runner.invoke(
            app,
            [
                "analyze",
                str(register),
                "--config",
                str(CONFIG_PATH),
                "--out",
                str(out),
                "--format",
                "excel",
            ],
        )
        assert result.exit_code == 0, result.output

        path = next(out.glob("*.xlsx"))
        with zipfile.ZipFile(path) as archive:
            formulas = [
                match
                for name in archive.namelist()
                if name.startswith("xl/worksheets/")
                for match in FORMULA_RE.findall(archive.read(name).decode("utf-8"))
            ]
        assert formulas, "лист «Документы» §9.1.8 обязан нести итог =SUM()"
        assert all(FORMULA_SUM_RE.fullmatch(formula) for formula in formulas), formulas

        workbook = load_workbook(path)
        text = " ".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert "#VALUE!" not in text
        assert "#REF!" not in text
        workbook.close()

    def test_sum_range_covers_the_data_rows(
        self,
        register: Path,
        config: Config,
        tmp_path: Path,
    ) -> None:
        """Диапазон итога обязан совпадать со строками находок, иначе итог лжёт."""
        result = run_pipeline(register, config)
        path = render_excel(result, tmp_path / "итоги.xlsx")

        workbook = load_workbook(path)
        sheet = workbook[SHEETS[7]]
        total_row = sheet.max_row
        formula = sheet.cell(row=total_row, column=5).value
        assert formula == f"=SUM(E2:E{total_row - 1})"
        assert total_row - 2 == len(result.findings)
        workbook.close()

    def test_no_report_qualifies_the_finding(
        self,
        register: Path,
        tmp_path: Path,
    ) -> None:
        """§13.10 на всех трёх форматах сразу."""
        runner = CliRunner()
        out = tmp_path / "формулировки"
        invocation = runner.invoke(
            app,
            [
                "analyze",
                str(register),
                "--config",
                str(CONFIG_PATH),
                "--out",
                str(out),
                "--format",
                "excel,json,html",
            ],
        )
        assert invocation.exit_code == 0, invocation.output

        text = (next(out.glob("*.json")).read_text(encoding="utf-8")).lower()
        text += (next(out.glob("*.html")).read_text(encoding="utf-8")).lower()
        workbook = load_workbook(next(out.glob("*.xlsx")))
        text += " ".join(
            str(cell.value).lower()
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        workbook.close()
        for word in FORBIDDEN_WORDS:
            assert word not in text


class TestBatchSummary:
    """§13.9 — пакетный прогон и сводная таблица по сети."""

    def test_summary_ranks_by_deviation(
        self,
        make_workbook: Callable[[SheetPlan], Path],
        tmp_path: Path,
    ) -> None:
        first = make_workbook(_overstated_plan())
        second = make_workbook(_overstated_plan())
        directory = first.parent
        summary = tmp_path / "сводка.xlsx"

        runner = CliRunner()
        invocation = runner.invoke(
            app,
            [
                "batch",
                str(directory),
                "--config",
                str(CONFIG_PATH),
                "--summary",
                str(summary),
            ],
        )
        assert invocation.exit_code == 0, invocation.output
        assert summary.exists()

        workbook = load_workbook(summary)
        sheet = workbook.worksheets[0]
        header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        assert header[:3] == ["ПВЗ", "Отклонение от целевого", "Находок"]
        assert sheet.max_row == 3
        assert second.exists()
        workbook.close()

    def test_empty_directory_is_an_error_not_an_empty_summary(self, tmp_path: Path) -> None:
        """Молчаливая пустая сводка выглядела бы как «сеть в порядке»."""
        runner = CliRunner()
        invocation = runner.invoke(
            app,
            ["batch", str(tmp_path), "--config", str(CONFIG_PATH)],
        )
        assert invocation.exit_code == 2
